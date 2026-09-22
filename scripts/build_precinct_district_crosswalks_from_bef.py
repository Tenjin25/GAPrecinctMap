"""Build VTD20-to-district weights from official Census/Georgia block equivalency files."""

from __future__ import annotations

import argparse
import csv
import io
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
import shapefile


def block_to_vtd(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        with archive.open("BlockAssign_ST13_GA_VTD.txt") as raw:
            reader = csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig"), delimiter="|")
            return {
                str(row["BLOCKID"]).strip(): f"13{str(row['COUNTYFP']).strip().zfill(3)}{str(row['DISTRICT']).strip()}"
                for row in reader
                if str(row.get("BLOCKID") or "").strip() and str(row.get("DISTRICT") or "").strip()
            }


def read_census_zip(path: Path, member: str, district_column: str) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        with archive.open(member) as raw:
            reader = csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig"), skipinitialspace=True)
            return {
                str(row["GEOID"]).strip(): str(row[district_column]).strip()
                for row in reader
                if str(row.get("GEOID") or "").strip() and str(row.get(district_column) or "").strip()
            }


def read_workbook(path: Path) -> dict[str, str]:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)
    return {
        str(block).strip(): str(district).strip().zfill(3)
        for block, district, *_ in rows
        if block is not None and district is not None
    }


def block_land_area(path: Path) -> dict[str, int]:
    """Read Census ALAND20 without retaining extracted shapefile components."""
    with tempfile.TemporaryDirectory() as temp_dir, zipfile.ZipFile(path) as archive:
        archive.extractall(temp_dir)
        shp = next(Path(temp_dir).glob("*.shp"))
        reader = shapefile.Reader(str(shp))
        result = {
            str(record["GEOID20"]): int(record["ALAND20"] or 0)
            for record in reader.iterRecords()
        }
        reader.close()
        return result


def build_rows(
    block_vtd: dict[str, str],
    block_district: dict[str, str],
    land_area: dict[str, int],
) -> list[dict[str, object]]:
    pairs: dict[tuple[str, str], int] = defaultdict(int)
    pair_area: dict[tuple[str, str], int] = defaultdict(int)
    totals: dict[str, int] = defaultdict(int)
    total_area: dict[str, int] = defaultdict(int)
    for block, district in block_district.items():
        vtd = block_vtd.get(block)
        if not vtd:
            continue
        pairs[(vtd, district)] += 1
        area = land_area.get(block, 0)
        pair_area[(vtd, district)] += area
        totals[vtd] += 1
        total_area[vtd] += area
    rows = []
    for (vtd, district), count in sorted(pairs.items()):
        weight = (pair_area[(vtd, district)] / total_area[vtd]) if total_area[vtd] else (count / totals[vtd])
        rows.append({
            "precinct_key": vtd,
            "district_num": str(int(district)) if district.isdigit() else district,
            "area_weight": f"{weight:.10f}",
            "land_area": pair_area[(vtd, district)],
            "precinct_land_area": total_area[vtd],
            "block_count": count,
            "precinct_block_count": totals[vtd],
        })
    return rows


def write(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "precinct_key", "district_num", "area_weight", "land_area", "precinct_land_area",
            "block_count", "precinct_block_count"
        ])
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {path} ({len(rows)} rows; {len({r['precinct_key'] for r in rows})} VTDs)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--blockassign", type=Path, default=Path("Data/BlockAssign_ST13_GA.zip"))
    parser.add_argument("--out-dir", type=Path, default=Path("Data/crosswalks_bef"))
    args = parser.parse_args()
    vtd = block_to_vtd(args.blockassign)
    area = block_land_area(Path("Data/tl_2020_13_tabblock20.zip"))
    inputs = [
        ("precinct_to_cd118.csv", read_census_zip(Path("Data/cd118.zip"), "13_GA_CD118.txt", "CDFP")),
        ("precinct_to_2022_state_house.csv", read_census_zip(Path("Data/sldl_2022.zip"), "13_GA_SLDL22.txt", "SLDLST")),
        ("precinct_to_2022_state_senate.csv", read_census_zip(Path("Data/sldu_2022.zip"), "13_GA_SLDU22.txt", "SLDUST")),
        ("precinct_to_cd119.csv", read_workbook(Path("Data/congress-2023-block-equiv.xlsx"))),
        ("precinct_to_2024_state_house.csv", read_workbook(Path("Data/house-2023-block-equiv.xlsx"))),
        ("precinct_to_2024_state_senate.csv", read_workbook(Path("Data/senate-2023-block-equivalency.xlsx"))),
    ]
    for name, assignments in inputs:
        write(args.out_dir / name, build_rows(vtd, assignments, area))


if __name__ == "__main__":
    main()
